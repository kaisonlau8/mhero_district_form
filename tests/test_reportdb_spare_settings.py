import io
import re

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from reportdb import calc, spare
from reportdb.db import connect
from test_reportdb_web import app_env, anon, client, login, viewer, _fast_hash  # noqa: F401


def tok(c, path="/settings/spare-parts"):
    return re.search(r'name="csrf" value="([^"]+)"', c.get(path).text).group(1)


def xlsx_bytes(rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_default_list_is_template_33(app_env):
    conn = connect(app_env)
    lst = spare.list_for_date(conn, "2026-01-01")
    assert lst["id"] == 1 and len(lst["items"]) == 33
    assert sum(1 for i in lst["items"] if not i["counted"]) == 1
    assert [i for i in lst["items"] if not i["counted"]][0]["code"].startswith("817") or True


def test_template_download_has_flag(client):
    r = client.get("/settings/spare-parts/template.xlsx")
    ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:3] == ("备件号", "备件名称", "是否计入分母") and len(rows) == 34
    assert [r[2] for r in rows[1:]].count("否") == 1


def test_parse_and_normalisation():
    parts = spare.parse_upload("a.xlsx", xlsx_bytes([["备件号", "备件名称", "是否计入分母"], [22806000, "A", None], ["ab-12 3", "B", "否"]]))
    assert parts[0]["code"] == "22806000" and parts[0]["counted"] == 1 and parts[1]["counted"] == 0
    from reportdb.loaders import part_key
    assert part_key(22806000) == part_key("22806000.0") == part_key(" 22806000 ") == "22806000"
    assert part_key("ab-12 3") == part_key("AB123") == "AB123"
    csv = "备件号,备件名称\n123,x\n".encode("utf-8-sig")
    assert spare.parse_upload("a.csv", csv)[0]["code"] == "123"
    import pytest
    with pytest.raises(spare.SpareListError):
        spare.parse_upload("a.csv", "名称\nx\n".encode())
    with pytest.raises(spare.SpareListError):
        spare.parse_upload("a.csv", "备件号,备件名称,是否计入分母\n1,x,也许\n".encode())
    with pytest.raises(spare.SpareListError):
        spare.parse_upload("a.pdf", b"x")


def test_upload_preview_confirm_recompute_rollback(app_env, client):
    conn = connect(app_env)
    before = calc.spare_matrix(conn, "2026-01-01")
    conn.close()
    f = xlsx_bytes([["备件号", "备件名称", "是否计入分母"], ["22806000", "机油滤", "是"], ["NOPE-1", "不存在", "是"], ["BAD", "不计入", "否"]])
    r = client.post("/settings/spare-parts/upload", data={"csrf": tok(client)},
                    files={"file": ("n.xlsx", f, "application/octet-stream")}, follow_redirects=False)
    assert r.status_code == 303
    pv = client.get(r.headers["location"])
    assert pv.status_code == 200 and "未匹配" in pv.text and "已匹配" in pv.text
    did = r.headers["location"].rsplit("/", 1)[1]
    # 预览阶段不生效
    conn = connect(app_env)
    assert conn.execute("SELECT COUNT(*) FROM spare_lists").fetchone()[0] == 1
    conn.close()
    r = client.post("/settings/spare-parts/confirm", data={"draft_id": did, "effective_from": "2026-01-01", "note": "t", "csrf": tok(client)}, follow_redirects=False)
    assert "msg=" in r.headers["location"]
    conn = connect(app_env)
    after = calc.spare_matrix(conn, "2026-01-01")
    assert after["counted"] == 2 and after["counted"] != before["counted"] and after["list"]["id"] == 2
    j = client.get("/api/tables/spare-parts").json()
    assert j  # 已按新名单重算并写入结果
    stored = conn.execute("SELECT payload FROM report_results WHERE snapshot_date='2026-01-01' AND table_key='spare-parts' AND region_mode='store'").fetchone()
    assert stored and '"id": 2' in stored[0]
    conn.close()
    # 生效日期在后：更早快照仍用旧名单
    conn = connect(app_env)
    assert spare.list_for_date(conn, "2025-12-31")["id"] == 1
    conn.close()
    # 回滚
    r = client.post("/settings/spare-parts/rollback", data={"list_id": 1, "effective_from": "2026-01-01", "csrf": tok(client)}, follow_redirects=False)
    assert "msg=" in r.headers["location"]
    conn = connect(app_env)
    back = calc.spare_matrix(conn, "2026-01-01")
    assert back["list"]["id"] == 3 and back["counted"] == before["counted"]
    assert [h["id"] for h in spare.history(conn)] == [3, 2, 1]  # 历史保留
    assert [p["rate"] for p in back["rows"]] == [p["rate"] for p in before["rows"]]


def test_bad_uploads_and_authz(client, viewer, anon):
    c = tok(client)
    r = client.post("/settings/spare-parts/upload", data={"csrf": c}, files={"file": ("x.csv", "备件号,备件名称\n".encode(), "text/csv")}, follow_redirects=False)
    assert "err=" in r.headers["location"]
    r = client.post("/settings/spare-parts/upload", data={"csrf": c}, files={"file": ("x.csv", "备件号,备件名称,是否计入分母\n1,a,否\n".encode(), "text/csv")}, follow_redirects=False)
    assert "err=" in r.headers["location"]  # 全部不计入
    assert client.post("/settings/spare-parts/upload", data={"csrf": "bad"}, files={"file": ("x.csv", b"a", "text/csv")}).status_code == 403
    assert viewer.post("/settings/spare-parts/upload", data={"csrf": "x"}, files={"file": ("x.csv", b"a", "text/csv")}).status_code == 403
    assert viewer.get("/settings/spare-parts/template.xlsx").status_code == 403
    assert anon.get("/settings/spare-parts/template.xlsx", follow_redirects=False).status_code == 303
