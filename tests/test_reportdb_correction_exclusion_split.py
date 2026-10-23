"""申请分两类：修正结果（按指标，否→是）与剔除（整车，批准后所有指标都不计）。含迁移、五个指标、页面。"""
import re
import sqlite3
from types import SimpleNamespace

import pytest

from reportdb import calc, db as dbm, exemptions as ex
from reportdb.ingest import ingest_datasets
from reportdb.db import connect
from test_reportdb_calc import D1, D2, row
from test_reportdb_exemptions import SPEC, SNAP1, U, counts, tok, web, env, load, vin_of  # noqa: F401

SHARED = "SHARED0000000001"   # D1，五个数据集里都有，结果都是「否」
SHARED2 = "SHARED0000000002"  # D1，五个数据集里都有，结果都是「是」


def load_shared(conn, snap):
    st = {r["name"]: r["code"] for r in conn.execute("SELECT code,name FROM stores")}
    data = {"stock": [], "task": [], "first-service": [], "second-service": [], "new-policy": [], "renewal": []}
    for ind, (ds, width, vi, di, hit, miss) in SPEC.items():
        def mk(vin, done):
            return row(width, {vi: vin, 1: "x", 2: st[D1], 3: D1, di: done})
        data[ds] = [mk(SHARED, miss), mk(SHARED2, hit), mk("OTHER00000000003", miss)]
    ingest_datasets(conn, snap, data, source="t", origin_renewal=["renewal"] * 3)


@pytest.fixture()
def shared(conn):
    conn.execute("UPDATE stores SET supervisor='张督导' WHERE name=?", (D1,))
    conn.commit()
    load_shared(conn, SNAP1)
    return SimpleNamespace(conn=conn, admin=U(conn, "admin", "admin", ""), sup=U(conn, "sup1", "supervisor", "张督导"))


def test_exclusion_is_vehicle_level_all_five_indicators(shared):
    c = shared.conn
    before = {i: counts(c, i, SNAP1)[D1.casefold()] for i in SPEC}
    assert all(v == [3, 1] for v in before.values())
    r = ex.request_exclusions(c, shared.sup, SNAP1, items=[(SHARED, "出口车")])
    assert len(r["created"]) == 1
    row_ = c.execute("SELECT indicator, kind, status FROM exemptions WHERE id=?", (r["created"][0],)).fetchone()
    assert (row_["indicator"], row_["kind"], row_["status"]) == ("all", "exclude", "pending")
    assert {i: counts(c, i, SNAP1)[D1.casefold()] for i in SPEC} == before, "待审批不生效"
    ex.review(c, shared.admin, r["created"], "approve", "同意", expires_at="2099-01-01")
    for i in SPEC:   # 同一辆车在五个指标里都被剔除：分母 -1（未完成的车，分子不变）
        assert counts(c, i, SNAP1)[D1.casefold()] == [2, 1], i
    assert c.execute("SELECT expires_at FROM exemptions WHERE id=?", (r["created"][0],)).fetchone()[0] is None   # 永久
    # 汇总表 / 日报用的引擎结果也同步
    sheet = calc.compute(c, SNAP1)["direct-stores"].blocks[0].rows[0]
    assert sheet[0] == D1
    # 新快照继续沿用
    load_shared(c, "2026-01-02")
    for i in SPEC:
        assert counts(c, i, "2026-01-02")[D1.casefold()] == [2, 1], i
    # 同一 VIN 不能重复申请
    again = ex.request_exclusions(c, shared.sup, SNAP1, items=[(SHARED, "x")])
    assert again["created"] == []
    ex.review(c, shared.admin, r["created"], "revoke", "回退")
    for i in SPEC:
        assert counts(c, i, SNAP1)[D1.casefold()] == [3, 1], i


def test_correction_is_per_indicator_only_and_only_for_no(shared):
    c = shared.conn
    for ind in SPEC:
        done = ex.request_exemptions(c, shared.sup, ind, SNAP1, vins=[SHARED2], category="其他", reason="x")
        assert done["created"] == [] and "无需修正" in done["skipped"][0][1], ind   # 结果已是「是」
        assert [m["vin"] for m in ex.candidates(c, ind, SNAP1, ex.scope_for(c, shared.sup))].count(SHARED2) == 0
    r = ex.request_exemptions(c, shared.sup, "second_service", SNAP1, vins=[SHARED], category="其他", reason="补录")
    ex.review(c, shared.admin, r["created"], "approve")
    assert counts(c, "second_service", SNAP1)[D1.casefold()] == [3, 2]
    for ind in SPEC:
        if ind != "second_service":
            assert counts(c, ind, SNAP1)[D1.casefold()] == [3, 1], "只改被修正的那个指标"
    # 之后剔除该车：修正的那条也不再计（整车不计）
    x = ex.request_exclusions(c, shared.sup, SNAP1, items=[(SHARED, "出口车")])
    ex.review(c, shared.admin, x["created"], "approve")
    assert counts(c, "second_service", SNAP1)[D1.casefold()] == [2, 1]
    # 已被剔除的车不再出现在修正结果列表，也不能再申请修正
    assert SHARED not in [m["vin_key"] if "vin_key" in m else m["vin"] for m in ex.candidates(c, "first_service", SNAP1, None)]
    assert ex.request_exemptions(c, shared.sup, "first_service", SNAP1, vins=[SHARED], category="其他", reason="x")["created"] == []
    # 两类互不通用
    with pytest.raises(ex.ExemptionError):
        ex.request_exemptions(c, shared.sup, "first_service", SNAP1, vins=[SHARED], reason="x", kind="exclude")


def test_vehicle_candidates_scope_dedupe(shared):
    c = shared.conn
    mine = ex.vehicle_candidates(c, SNAP1, ex.scope_for(c, shared.sup))
    assert len({m["vin_key"] for m in mine}) == len(mine) == 3   # 五个数据集里同一辆车只出现一次
    assert all(len(m["inds"]) == 5 for m in mine if m["vin_key"] in (SHARED, SHARED2))
    ghost = U(c, "sup9", "supervisor", "没有门店")
    assert ex.vehicle_candidates(c, SNAP1, ex.scope_for(c, ghost)) == []
    out = ex.request_exclusions(c, ghost, SNAP1, items=[(SHARED, "x")])
    assert out["created"] == [] and "范围" in out["skipped"][0][1]


# ------------------------------------------------------------------ 页面
def test_two_separate_pages_and_flow(web):
    sup, adm = web["sup1"], web["admin"]
    nav = sup.get("/").text
    assert "申请修正结果" in nav and "申请剔除" in nav and "/exclusions/new" in nav and "豁免" not in nav
    ex_page = sup.get("/exclusions/new").text
    assert "整辆车" in ex_page and "我的申请明细（剔除）" in ex_page and "提交申请" in ex_page
    assert "证据说明" not in ex_page and 'name="vins"' not in ex_page and "指标</option>" not in ex_page
    assert 'name="kind"' not in ex_page
    picks = re.findall(r'name="xreason_([^"]+)"', ex_page)
    assert "xpick_" not in ex_page and "是否申请" not in ex_page and ex_page.count('class="ant-btn ant-btn-primary" type="submit">提交申请') == len(picks)
    assert len(picks) == 10, picks   # 本店 5 个指标 × 2 辆车（各指标 VIN 不同），别人门店/不在门店表的没有
    v = vin_of("first_service", "A2")   # 结果已是「是」的车：剔除照样可以申请
    assert v in picks
    t = tok(sup, "/exclusions/new")
    # 没填理由：不提交；一次多辆：整体拒绝
    r = sup.post("/exclusions/request", data={"date": SNAP1, "csrf": t, f"xreason_{v}": ""}, follow_redirects=False)
    assert "err=" in r.headers["location"]
    v2 = next(p for p in picks if p != v)
    r = sup.post("/exclusions/request", data={"date": SNAP1, "csrf": t, f"xreason_{v}": "出口车", f"xreason_{v2}": "出口车"}, follow_redirects=False)
    assert "err=" in r.headers["location"] and "1+%E8%BE%86" in r.headers["location"]
    assert sqlite3.connect(web["db"]).execute("SELECT COUNT(*) FROM exemptions").fetchone()[0] == 0
    assert sup.post("/exclusions/request", data={"date": SNAP1, "csrf": "bad", f"xreason_{v}": "x"}).status_code == 403
    r = sup.post("/exclusions/request", data={"date": SNAP1, "csrf": t, f"xreason_{v}": "出口车，DMS 没标"}, follow_redirects=False)
    assert "msg=" in r.headers["location"] and "/exclusions/new" in r.headers["location"]
    c = sqlite3.connect(web["db"]); c.row_factory = sqlite3.Row
    rows = c.execute("SELECT indicator, kind, status FROM exemptions").fetchall()
    assert [(x["indicator"], x["kind"], x["status"]) for x in rows] == [("all", "exclude", "pending")]
    assert v in sup.get("/exclusions/new").text and "出口车，DMS 没标" in sup.get("/exclusions/new").text   # 我的申请明细（剔除）
    mine = sup.get("/exemptions/mine").text
    assert "剔除" in mine and "全部指标（整车）" in mine
    # 修正结果页不显示剔除申请；提交修正结果是另一类
    assert "出口车，DMS 没标" not in sup.get("/exemptions/new?indicator=first_service").text
    # 管理员各类分别审批
    assert adm.get("/admin/exemptions?kind=exempt").text.count('name="ids"') == 0
    page = adm.get("/admin/exemptions?kind=exclude")
    assert "剔除审批" in page.text and "全部指标（整车）" in page.text
    eid = re.search(r'name="ids" value="(\d+)"', page.text).group(1)
    r = adm.post("/admin/exemptions/review", data={"action": "approve", "ids": eid, "csrf": tok(adm, "/admin/exemptions")}, follow_redirects=False)
    assert "msg=" in r.headers["location"]
    # 批准后：该车在首保明细标「已剔除」，不再出现在两个申请页，首保分母 -1
    assert '<span class="ant-tag ant-tag-warning">已剔除</span>' in sup.get("/data/first-service").text
    assert f'name="xreason_{v}"' not in sup.get("/exclusions/new").text
    assert adm.get("/api/tables/direct-stores").json()["blocks"][0]["rows"][0][1] == 1
    assert "我的申请明细" in sup.get("/exemptions/new?indicator=first_service").text
    d = adm.get(f"/admin/exemptions/{eid}").text
    assert "剔除" in d and "所有指标" in d
    from openpyxl import load_workbook
    import io
    assert load_workbook(io.BytesIO(adm.get("/admin/exemptions.xlsx").content)).active.max_row == 2


def test_view_only_user_cannot_use_exclusion_pages(web):
    assert web["view1"].get("/exclusions/new").status_code == 403
    assert web["view1"].post("/exclusions/request", data={"csrf": "x"}).status_code in (403, 401)


# ------------------------------------------------------------------ 迁移（v4 → v5）
def _v4_db(tmp_path):
    p = tmp_path / "old.sqlite3"
    c = connect(p)
    for i, m in enumerate(dbm.MIGRATIONS[:4]):
        c.executescript(m) if isinstance(m, str) else m(c)
        c.execute(f"PRAGMA user_version={i + 1}")
        c.commit()
    dbm.seed_if_empty(c)
    dbm.ensure_default_spare_list(c)
    return p, c


def _ins(c, vin, ind, kind, status, eff="2026-01-02", exp=None, rid=None):
    cur = c.execute("INSERT INTO exemptions(vin,vin_key,indicator,reason_category,reason,requested_by,requested_at,status,kind,effective_from,expires_at) "
                    "VALUES (?,?,?,?,?,?,?,?,?,?,?)", (vin, vin.upper(), ind, "其他", "r", "u", "2026-01-01 00:00:00", status, kind, eff, exp))
    return cur.lastrowid


def test_migration_v5_per_indicator_exclusions_become_vehicle_level(tmp_path):
    p, c = _v4_db(tmp_path)
    ex1 = _ins(c, "AAA", "second_service", "exempt", "approved")
    a = _ins(c, "BBB", "task", "exclude", "approved", eff="2026-01-05", exp="2027-01-01")
    b = _ins(c, "BBB", "first_service", "exclude", "approved", eff="2026-01-03")
    d = _ins(c, "BBB", "renewal", "exclude", "pending", eff=None)
    e = _ins(c, "CCC", "task", "exclude", "pending", eff=None)
    f = _ins(c, "DDD", "task", "exclude", "rejected", eff=None)
    g = _ins(c, "EEE", "new_policy", "exclude", "pending", eff=None)
    h = _ins(c, "EEE", "renewal", "exclude", "pending", eff=None)
    c.commit()
    assert dbm.schema_version(c) == 4
    assert dbm.migrate(c) == len(dbm.MIGRATIONS)
    st = {r["id"]: dict(r) for r in c.execute("SELECT * FROM exemptions")}
    assert st[ex1]["kind"] == "exempt" and st[ex1]["indicator"] == "second_service" and st[ex1]["status"] == "approved"   # 豁免 → 修正结果，原样
    assert all(st[i]["indicator"] == "all" for i in (a, b, d, e, f, g, h))
    # BBB：留最早的已批准（#a，生效日期取最早 01-03，永久），其余合并
    assert st[a]["status"] == "approved" and st[a]["effective_from"] == "2026-01-03" and st[a]["expires_at"] is None
    assert st[b]["status"] == "revoked" and st[d]["status"] == "withdrawn"
    assert st[e]["status"] == "pending" and st[f]["status"] == "rejected"
    assert st[g]["status"] == "pending" and st[h]["status"] == "withdrawn"
    live = c.execute("SELECT vin_key, COUNT(*) n FROM exemptions WHERE kind='exclude' AND status IN ('pending','approved') GROUP BY vin_key").fetchall()
    assert {r["vin_key"]: r["n"] for r in live} == {"BBB": 1, "CCC": 1, "EEE": 1}
    assert c.execute("SELECT COUNT(*) FROM exemption_log WHERE actor='migration'").fetchone()[0] >= 3
    assert (tmp_path / "old.sqlite3.pre-v5.bak").exists()
    bak = sqlite3.connect(str(tmp_path / "old.sqlite3.pre-v5.bak"))
    assert bak.execute("PRAGMA user_version").fetchone()[0] == 4
    assert bak.execute("SELECT COUNT(*) FROM exemptions WHERE kind='exclude' AND indicator<>'all'").fetchone()[0] == 7
    # 幂等：再迁移一次什么也不变
    snap = [tuple(r) for r in c.execute("SELECT * FROM exemptions ORDER BY id")]
    nlog = c.execute("SELECT COUNT(*) FROM exemption_log").fetchone()[0]
    assert dbm.migrate(c) == len(dbm.MIGRATIONS)
    assert [tuple(r) for r in c.execute("SELECT * FROM exemptions ORDER BY id")] == snap
    assert c.execute("SELECT COUNT(*) FROM exemption_log").fetchone()[0] == nlog


def test_migration_v5_recomputes_snapshots_when_approved_exclusion_exists(tmp_path):
    p, c = _v4_db(tmp_path)
    c.execute("UPDATE stores SET supervisor='张督导' WHERE name=?", (D1,)); c.commit()
    load_shared(c, SNAP1)
    def payload():
        return c.execute("SELECT payload FROM report_results WHERE snapshot_date=? AND table_key='direct-stores' AND region_mode='store'", (SNAP1,)).fetchone()[0]
    before = payload()   # 迁移前存下的汇总（没有任何剔除）
    # 旧：SHARED 只在「首保」上被剔除（按指标）；迁移后所有指标都剔除并重算汇总
    _ins(c, SHARED, "first_service", "exclude", "approved", eff=SNAP1)
    c.commit()
    assert payload() == before
    dbm.migrate(c)
    for i in SPEC:
        assert counts(c, i, SNAP1)[D1.casefold()] == [2, 1], i
    assert payload() != before   # 汇总结果已按新口径重算
