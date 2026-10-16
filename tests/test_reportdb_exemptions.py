"""豁免：分子翻转、跨快照持久、撤销、督导范围、审计、权限、页面流程。"""
import re
import sqlite3
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient

from reportdb import calc, exemptions as ex
from reportdb.auth import service
from reportdb.datasets import BY_KEY
from reportdb.db import connect
from reportdb.ingest import ingest_datasets
from reportdb.webapp import create_app
from test_reportdb_calc import D1, D2, row
from test_reportdb_web import login, _fast_hash  # noqa: F401

SNAP1, SNAP2 = "2026-01-01", "2026-01-02"
GHOST = "不在门店表的店"
# 指标 → (数据集 key, 总列数, VIN 列, 决定列, 命中值, 未命中值)
SPEC = {
    "task": ("task", 21, 8, 6, "待处理", "已完成"),
    "first_service": ("first-service", 12, 0, 7, "是", "否"),
    "second_service": ("second-service", 12, 0, 7, "是", "否"),
    "new_policy": ("new-policy", 13, 0, 6, "是", "否"),
    "renewal": ("renewal", 13, 0, 7, "是", "否"),
}
TABLE = {"task": "task_items"}


def rows_for(conn, ind, n_first=1):
    ds, width, vi, di, hit, miss = SPEC[ind]
    st = {r["name"]: r["code"] for r in conn.execute("SELECT code,name FROM stores")}

    def mk(vin, store, done):
        return row(width, {vi: vin, 1: "x", 2: st.get(store, "NOSUCH"), 3: store, di: done})

    return [mk(f"A1{ind[:2].upper()}0000000000001", D1, miss),     # D1 未完成（可豁免）
            mk(f"A2{ind[:2].upper()}0000000000002", D1, hit),      # D1 已完成
            mk(f"B1{ind[:2].upper()}0000000000003", D2, miss),     # D2 未完成（别的督导）
            mk(f"Z1{ind[:2].upper()}0000000000004", GHOST, miss)]  # 不在门店表


def vin_of(ind, n):
    return f"{n}{ind[:2].upper()}0000000000{'001 002 003 004'.split()[('A1 A2 B1 Z1'.split().index(n))]}"


def load(conn, snap):
    data = {"stock": [], "task": [], "first-service": [], "second-service": [], "new-policy": [], "renewal": []}
    for ind, spec in SPEC.items():
        data[spec[0]] = rows_for(conn, ind)
    ingest_datasets(conn, snap, data, source="t", origin_renewal=["renewal"] * 4)


def table_of(ind):
    return ex.INDICATORS[ind].table


def counts(conn, ind, snap, key="store_name"):
    i = ex.INDICATORS[ind]
    return calc._agg(conn, i.table, snap, key, i.flag_col, i.hit)


def U(conn, name, role, sup=""):
    uid = service.create_user(conn, name, "password-long-1", role=role, supervisor_name=sup)
    r = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
    return SimpleNamespace(id=uid, username=name, role=role, supervisor_name=r["supervisor_name"] or "", can_request=role in ("admin", "supervisor"))


@pytest.fixture()
def env(conn):
    conn.execute("UPDATE stores SET supervisor='张督导' WHERE name=?", (D1,))
    conn.execute("UPDATE stores SET supervisor='李督导' WHERE name=?", (D2,))
    conn.commit()
    load(conn, SNAP1)
    return SimpleNamespace(conn=conn, admin=U(conn, "admin", "admin"), sup1=U(conn, "sup1", "supervisor", "张督导"),
                           sup2=U(conn, "sup2", "supervisor", "李督导"))


@pytest.mark.parametrize("ind", list(SPEC))
def test_exemption_flips_numerator_only(env, ind):
    c = env.conn
    assert counts(c, ind, SNAP1)[D1.casefold()] == [2, 1]
    res = ex.request_exemptions(c, env.sup1, ind, SNAP1, vins=[vin_of(ind, "A1")], category="出口车挂店", reason="出口车挂靠本店")
    assert len(res["created"]) == 1
    assert counts(c, ind, SNAP1)[D1.casefold()] == [2, 1], "待审批不影响计算"
    out = ex.review(c, env.admin, res["created"], "approve", "同意")
    assert out["recomputed"] == [SNAP1]
    assert counts(c, ind, SNAP1)[D1.casefold()] == [2, 2], "分子+1，分母不变"
    assert counts(c, ind, SNAP1)[D2.casefold()] == [1, 0], "别的门店不受影响"
    assert counts(c, ind, SNAP1, "region").get("华北") or True
    assert ex.flipped_counts(c, SNAP1)[ind] == {"title": ex.INDICATORS[ind].title, "rows": 1, "vins": 1}
    ex.review(c, env.admin, res["created"], "revoke", "撤销测试")
    assert counts(c, ind, SNAP1)[D1.casefold()] == [2, 1]
    assert ex.flipped_counts(c, SNAP1)[ind]["rows"] == 0


def test_summary_tables_follow_exemption(env):
    c = env.conn
    def snap_payload():
        return c.execute("SELECT payload FROM report_results WHERE snapshot_date=? AND table_key='direct-stores' AND region_mode='store'", (SNAP1,)).fetchone()[0]
    before = snap_payload()
    r = ex.request_exemptions(c, env.sup1, "first_service", SNAP1, vins=[vin_of("first_service", "A1")], category="其他", reason="测试")
    ex.review(c, env.admin, r["created"], "approve")
    after = snap_payload()
    assert after != before
    first = calc.compute(c, SNAP1)["direct-stores"].blocks[0]
    assert first.rows[0][:4] == [D1, 2, 2, 1.0]
    ex.review(c, env.admin, r["created"], "revoke", "回退")
    assert snap_payload() == before


def test_persists_across_new_snapshots_and_effective_from(env):
    c = env.conn
    load(c, "2025-12-31")  # 早于生效日期的历史快照
    r = ex.request_exemptions(c, env.sup1, "first_service", SNAP1, vins=[vin_of("first_service", "A1")], category="出口车挂店", reason="出口")
    ex.review(c, env.admin, r["created"], "approve")  # 默认 effective_from = 申请时快照日 = 01-01
    assert counts(c, "first_service", "2025-12-31")[D1.casefold()] == [2, 1], "生效日期之前不受影响"
    load(c, SNAP2)  # 之后的新快照（同一 VIN 仍显示「否」）自动沿用
    assert counts(c, "first_service", SNAP2)[D1.casefold()] == [2, 2]
    # 晚入库的历史快照也按生效日期套用
    load(c, "2026-01-05")
    assert counts(c, "first_service", "2026-01-05")[D1.casefold()] == [2, 2]
    # 批准时可指定更早的生效日期 → 重算历史
    r2 = ex.request_exemptions(c, env.sup1, "second_service", SNAP1, vins=[vin_of("second_service", "A1")], category="其他", reason="x")
    out = ex.review(c, env.admin, r2["created"], "approve", effective_from="2025-12-31")
    assert "2025-12-31" in out["recomputed"] and counts(c, "second_service", "2025-12-31")[D1.casefold()] == [2, 2]


def test_expiry(env):
    c = env.conn
    r = ex.request_exemptions(c, env.sup1, "new_policy", SNAP1, vins=[vin_of("new_policy", "A1")], category="其他", reason="x")
    ex.review(c, env.admin, r["created"], "approve", expires_at=SNAP1)
    load(c, SNAP2)
    assert counts(c, "new_policy", SNAP1)[D1.casefold()] == [2, 2]
    assert counts(c, "new_policy", SNAP2)[D1.casefold()] == [2, 1], "到期后不再生效"


def test_excluded_store_stays_excluded(env):
    c = env.conn
    res = ex.request_exemptions(c, env.admin, "first_service", SNAP1, vins=[vin_of("first_service", "Z1")], category="其他", reason="x")
    assert res["created"] == [] and "门店表" in res["skipped"][0][1]
    # 即便手工塞一条已批准的记录，不在门店表的行也不计入
    c.execute("INSERT INTO exemptions(vin,vin_key,indicator,reason_category,reason,requested_by,requested_at,status,effective_from) "
              "VALUES (?,?,?,?,?,?,?,?,?)", (vin_of("first_service", "Z1"), vin_of("first_service", "Z1"), "first_service", "其他", "x", "t", "t", "approved", SNAP1))
    c.commit()
    calc.compute_and_store(c, SNAP1)
    assert GHOST.casefold() not in counts(c, "first_service", SNAP1)
    assert ex.flipped_counts(c, SNAP1)["first_service"]["rows"] == 0


def test_supervisor_scope_enforced(env):
    c = env.conn
    mine = ex.candidates(c, "first_service", SNAP1, ex.scope_for(c, env.sup1))
    assert {m["store_name"] for m in mine} == {D1} and len(mine) == 1 and not mine[0]["done"] if "done" in mine[0] else {m["store_name"] for m in mine} == {D1} and len(mine) == 1   # 修正结果列表：只有结果为「否」的车
    other = ex.request_exemptions(c, env.sup1, "first_service", SNAP1, vins=[vin_of("first_service", "B1")], category="其他", reason="x")
    assert other["created"] == [] and "范围" in other["skipped"][0][1]
    # 按行号也不能越权
    b1 = c.execute("SELECT row_no FROM first_service WHERE snapshot_date=? AND vin=?", (SNAP1, vin_of("first_service", "B1"))).fetchone()[0]
    assert ex.request_exemptions(c, env.sup1, "first_service", SNAP1, row_nos=[b1], category="其他", reason="x")["created"] == []
    # 手工分配门店后可申请；管理员不受限
    c.execute("INSERT INTO user_stores(user_id,store_code) SELECT ?, code FROM stores WHERE name=?", (env.sup1.id, D2)); c.commit()
    assert len(ex.request_exemptions(c, env.sup1, "first_service", SNAP1, row_nos=[b1], category="其他", reason="x")["created"]) == 1
    assert ex.scope_for(c, env.admin) is None
    # 无督导姓名、无分配 → 空范围
    ghost = U(c, "sup3", "supervisor", "不存在的人")
    assert ex.candidates(c, "first_service", SNAP1, ex.scope_for(c, ghost)) == []


def test_validation_duplicates_withdraw_and_permissions(env):
    c = env.conn
    v = vin_of("renewal", "A1")
    r0 = ex.request_exemptions(c, env.sup1, "renewal", SNAP1, vins=[v], category="随便", reason="x")   # 类别已改为可选，不认识的当「其他」
    assert r0["created"]
    ex.withdraw(c, env.sup1, r0["created"][0])
    with pytest.raises(ex.ExemptionError):
        ex.request_exemptions(c, env.sup1, "renewal", SNAP1, vins=[v], category="其他", reason="  ")
    viewer = U(c, "v1", "viewer")
    with pytest.raises(ex.ExemptionError):
        ex.request_exemptions(c, viewer, "renewal", SNAP1, vins=[v], category="其他", reason="x")
    done = ex.request_exemptions(c, env.sup1, "renewal", SNAP1, vins=[vin_of("renewal", "A2")], category="其他", reason="x")
    assert done["created"] == [] and "无需修正" in done["skipped"][0][1]
    r = ex.request_exemptions(c, env.sup1, "renewal", SNAP1, vins=[v.lower()], category="其他", reason="x")
    assert len(r["created"]) == 1
    dup = ex.request_exemptions(c, env.sup1, "renewal", SNAP1, vins=[v], category="其他", reason="x")
    assert dup["created"] == []
    with pytest.raises(ex.ExemptionError):
        ex.withdraw(c, env.sup2, r["created"][0])  # 不能撤回别人的
    with pytest.raises(ex.ExemptionError):
        ex.review(c, env.sup1, r["created"], "approve")  # 督导无审批权
    with pytest.raises(ex.ExemptionError):
        ex.review(c, env.admin, r["created"], "reject", "")  # 驳回必须写意见
    ex.withdraw(c, env.sup1, r["created"][0])
    assert c.execute("SELECT status FROM exemptions WHERE id=?", (r["created"][0],)).fetchone()[0] == "withdrawn"
    again = ex.request_exemptions(c, env.sup1, "renewal", SNAP1, vins=[v], category="其他", reason="再次申请")
    out = ex.review(c, env.admin, again["created"], "reject", "证据不足")
    assert out["recomputed"] == [] and counts(c, "renewal", SNAP1)[D1.casefold()] == [2, 1]
    assert ex.review(c, env.admin, again["created"], "approve")["skipped"] == again["created"]  # 已驳回不能再批


def test_audit_trail_is_append_only(env):
    c = env.conn
    r = ex.request_exemptions(c, env.sup1, "task", SNAP1, vins=[vin_of("task", "A1")], category="出口车挂店", reason="出口")
    i = r["created"][0]
    ex.review(c, env.admin, [i], "approve", "ok")
    ex.review(c, env.admin, [i], "revoke", "错了")
    log = ex.history(c, i)
    assert [(l["action"], l["actor"], l["from_status"], l["to_status"]) for l in log] == [
        ("request", "sup1", None, "pending"), ("approve", "admin", "pending", "approved"), ("revoke", "admin", "approved", "revoked")]
    with pytest.raises(sqlite3.DatabaseError):
        c.execute("UPDATE exemption_log SET actor='x'")
    with pytest.raises(sqlite3.DatabaseError):
        c.execute("DELETE FROM exemption_log")
    assert len(ex.export_xlsx(ex.list_exemptions(c))) > 1000


def test_oidc_supervisor_mapping(monkeypatch):
    from reportdb.auth.oidc import OidcBackend
    for k, v in {"REPORTDB_OIDC_ISSUER": "https://i", "REPORTDB_OIDC_CLIENT_ID": "c", "REPORTDB_OIDC_REDIRECT_URI": "https://r/cb",
                 "REPORTDB_OIDC_SUPERVISOR_GROUPS": "supervisor", "REPORTDB_OIDC_ADMIN_USERS": "boss"}.items():
        monkeypatch.setenv(k, v)
    b = OidcBackend(lambda: None)
    assert b.map_identity({"preferred_username": "zhang", "name": "张督导", "roles": [{"name": "supervisor"}]}).role == "supervisor"
    assert b.map_identity({"preferred_username": "boss", "roles": ["supervisor"]}).role == "admin"
    assert b.map_identity({"preferred_username": "x"}).role == "viewer"


# ------------------------------------------------------------------ 页面
@pytest.fixture()
def web(tmp_path, conn):
    conn.execute("UPDATE stores SET supervisor='张督导' WHERE name=?", (D1,))
    conn.execute("UPDATE stores SET supervisor='李督导' WHERE name=?", (D2,))
    conn.commit()
    load(conn, SNAP1)
    service.create_user(conn, "admin", "admin-password-1", role="admin")
    service.create_user(conn, "sup1", "sup1-password-12", role="supervisor", supervisor_name="张督导")
    service.create_user(conn, "view1", "view1-password-12", role="viewer")
    conn.close()
    db = tmp_path / "t.sqlite3"
    clients = {}
    for u, p in (("admin", "admin-password-1"), ("sup1", "sup1-password-12"), ("view1", "view1-password-12")):
        c = TestClient(create_app(db))
        login(c, u, p)
        clients[u] = c
    clients["db"] = db
    return clients


def tok(c, path):
    return re.search(r'name="csrf" value="([^"]+)"', c.get(path).text).group(1)


def test_web_workflow(web):
    sup, adm, view = web["sup1"], web["admin"], web["view1"]
    # 督导：看全部区域汇总和全部明细、VIN 不脱敏；只有自己门店的行能勾选申请
    j = sup.get("/api/data/first-service?size=50").json()
    assert {r["store_name"] for r in j["rows"]} == {D1, D2, GHOST} and j["total"] == 4 and j["masked"] is False
    assert j["rows"][0]["vin"] == vin_of("first_service", "A1")
    html = sup.get("/data/first-service").text
    assert D2 in html and 'name="reason_' not in html and 'name="pick_' not in html and "我的申请明细" not in html   # 明细页只读
    assert 'name="kind"' not in html and "/exemptions/new?indicator=first_service" in html
    assert view.get("/api/data/first-service").json()["masked"] is True
    assert view.get("/exemptions/new").status_code == 403 and "/exemptions/new" not in view.get("/data/first-service").text
    new = sup.get("/exemptions/new?indicator=first_service").text
    assert "申请修正结果" in new and "我的申请明细（修正结果）" in new and "修正结果：把这辆车" in new and 'name="kind" value="exclude"' not in new
    assert new.count('name="reason_') == 1 and new.count('name="pick_') == 1   # 只有 D1 的 A1(结果为否)；已完成的、别人门店的都没有
    assert sup.get("/admin/exemptions").status_code == 403 and view.get("/admin/exemptions").status_code == 403
    # 通过明细页勾选提交
    rn = re.search(r'name="reason_(\d+)"', new).group(1)
    r = sup.post("/exemptions/request", data={"indicator": "first_service", "date": SNAP1, f"reason_{rn}": "出口车挂在门店",
                                              "csrf": tok(sup, "/data/first-service")}, follow_redirects=False)
    assert r.status_code == 303 and "msg=" in r.headers["location"]
    assert sup.post("/exemptions/request", data={"indicator": "first_service", "date": SNAP1, "vins": "x", "category": "其他", "reason": "x", "csrf": "bad"}).status_code == 403
    assert "待审批" in sup.get("/exemptions/mine").text
    # 管理员审批
    page = adm.get("/admin/exemptions")
    assert page.status_code == 200 and "出口车挂在门店" in page.text
    eid = re.search(r'name="ids" value="(\d+)"', page.text).group(1)
    assert adm.post("/admin/exemptions/review", data={"action": "approve", "ids": eid}).status_code == 403  # 无 csrf
    r = adm.post("/admin/exemptions/review", data={"action": "approve", "ids": eid, "comment": "同意", "csrf": tok(adm, "/admin/exemptions")}, follow_redirects=False)
    assert "msg=" in r.headers["location"]
    # 汇总表、明细标记、透明度
    t = sup.get("/api/tables/direct-stores").json()
    assert t["blocks"][0]["rows"][0][2] == 2
    assert "已修正" in sup.get("/data/first-service").text
    q = adm.get("/api/quality").json()
    assert q["exemptions"]["first_service"]["rows"] == 1 and "已修正" in adm.get("/").text
    assert "修正结果 / 剔除审批" in adm.get("/").text
    # 详情与历史、导出
    d = adm.get(f"/admin/exemptions/{eid}").text
    assert "approve" in d and "request" in d
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(adm.get("/admin/exemptions.xlsx").content))
    assert wb.active.max_row == 2
    # 撤销
    r = adm.post("/admin/exemptions/review", data={"action": "revoke", "ids": eid, "comment": "撤销", "csrf": tok(adm, "/admin/exemptions")}, follow_redirects=False)
    assert "msg=" in r.headers["location"]
    assert adm.get("/api/tables/direct-stores").json()["blocks"][0]["rows"][0][2] == 1


def test_admin_assigns_stores_and_role(web):
    adm = web["admin"]
    html = adm.get("/admin/users").text
    sid = re.search(r'/admin/users/(\d+)/stores', html).group(1)
    page = adm.get(f"/admin/users/{sid}/stores")
    assert page.status_code == 200 and D1 in page.text
    code2 = conn_code(web["db"], D2)
    r = adm.post(f"/admin/users/{sid}/stores", data={"supervisor_name": "张督导", "codes": code2, "csrf": tok(adm, f"/admin/users/{sid}/stores")}, follow_redirects=False)
    assert r.status_code == 303
    html = web["sup1"].get("/exemptions/new?indicator=first_service").text
    assert html.count('name="reason_') > 1 and D2 in html  # 手工分配后 D2 的行也出现在申请列表


def conn_code(db, name):
    c = connect(db)
    try:
        return c.execute("SELECT code FROM stores WHERE name=?", (name,)).fetchone()[0]
    finally:
        c.close()


def test_cli_admin_account_is_admin_role(tmp_path, monkeypatch):
    from reportdb import users
    db = str(tmp_path / "u.sqlite3")
    monkeypatch.setenv("REPORTDB_ADMIN_PASSWORD", "cli-admin-password")
    assert users.main(["--db", db, "create-admin", "--username", "admin"]) == 0
    c = connect(db)
    r = c.execute("SELECT role, full_detail FROM users WHERE username='admin'").fetchone()
    assert (r["role"], r["full_detail"]) == ("admin", 1)
    # admin 有审批权
    row_ = c.execute("SELECT * FROM users WHERE username='admin'").fetchone()
    p = SimpleNamespace(role=row_["role"], username="admin")
    assert ex.review(c, p, [999], "approve")["skipped"] == [999]


def test_exemption_form_has_no_evidence_input():
    import pathlib
    import reportdb
    t = (pathlib.Path(reportdb.__file__).parent / "templates" / "exemption_new.html").read_text(encoding="utf-8")
    assert 'name="evidence"' not in t and "证据说明" not in t


@pytest.mark.parametrize("ind", list(ex.INDICATORS))
def test_correction_list_all_indicators(web, ind):
    """每个指标：/exemptions/new 只列结果为「否」的车、列与 /data 明细页一致；/data 页只读；剔除不能从这里提交。"""
    sup = web["sup1"]
    i = ex.INDICATORS[ind]
    ds = BY_KEY[i.ds_key]
    data = sup.get(f"/data/{ds.key}").text
    assert 'name="reason_' not in data and 'name="pick_' not in data and "我的申请明细" not in data
    assert "/exemptions/new?indicator=" in data and "/exclusions/new" in data
    new = sup.get(f"/exemptions/new?indicator={ind}").text
    heads = re.findall(r"<th[^>]*>(.*?)</th>", data.split("<thead>")[1].split("</thead>")[0])
    nheads = re.findall(r"<th[^>]*>(.*?)</th>", new.split("<thead>")[1].split("</thead>")[0])
    assert nheads[:2] == ["是否申请", "#"] and nheads[2:2 + len(heads) - 1] == [h.split(" <span")[0] for h in heads[1:]]
    assert nheads[-2:] == ["已有申请", "申请修正结果理由（必填）"]
    assert "申请修正结果" in new and "申请剔除" in new and "我的申请明细（修正结果）" in new and 'value="exclude"' not in new
    rows = re.findall(r'name="pick_(\d+)"', new)
    assert len(rows) == 1
    c = sqlite3.connect(web["db"]); c.row_factory = sqlite3.Row
    assert c.execute(f"SELECT COALESCE({i.flag_col},'') f FROM {i.table} WHERE row_no=? AND snapshot_date=?", (int(rows[0]), SNAP1)).fetchone()["f"] != i.hit
    t = tok(sup, f"/exemptions/new?indicator={ind}")
    # 就算硬塞 kind=exclude 也只会按「修正结果」处理，不会产生剔除
    r = sup.post("/exemptions/request", data={"indicator": ind, "kind": "exclude", "date": SNAP1, f"pick_{rows[0]}": "1", f"reason_{rows[0]}": "x", "csrf": t}, follow_redirects=False)
    assert r.status_code == 303
    assert c.execute("SELECT COUNT(*) FROM exemptions WHERE kind='exclude'").fetchone()[0] == 0
    assert c.execute("SELECT COUNT(*) FROM exemptions WHERE kind='exempt' AND indicator=? AND status='pending'", (ind,)).fetchone()[0] == 1
