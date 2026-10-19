import io
import re

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from reportdb.auth import service
from reportdb.ingest import ingest_datasets
from reportdb.webapp import create_app
from test_reportdb_calc import make_data

PAGES = ["/", "/tables/summary", "/tables/region-ranking", "/tables/direct-stores", "/tables/direct-ranking",
         "/tables/spare-parts", "/data/stock", "/data/task", "/data/first-service", "/data/second-service",
         "/data/new-policy", "/data/renewal"]


@pytest.fixture(autouse=True)
def _fast_hash(monkeypatch):
    monkeypatch.setenv("REPORTDB_PBKDF2_ITERS", "1000")


def login(client, username, password, expect=303):
    r = client.get("/login")
    csrf = re.search(r'name="csrf" value="([^"]+)"', r.text).group(1)
    r = client.post("/login", data={"username": username, "password": password, "csrf": csrf, "next": "/"},
                    follow_redirects=False)
    assert r.status_code == expect, r.text[:200]
    return r


@pytest.fixture()
def app_env(tmp_path, conn):
    data, origin = make_data(conn)
    data["task"][0][8] = "LDP29H925TM016610"   # 完整 VIN，用来测脱敏
    data["task"][0][9] = "王晓华"
    ingest_datasets(conn, "2026-01-01", data, source="t", origin_renewal=origin)
    service.create_user(conn, "boss", "boss-password-1", role="admin")
    service.create_user(conn, "viewer1", "viewer-password-1", role="viewer")
    service.create_user(conn, "auditor", "auditor-password-1", role="viewer", full_detail=True)
    conn.close()
    return tmp_path / "t.sqlite3"


@pytest.fixture()
def client(app_env):
    """已用管理员登录的客户端。"""
    c = TestClient(create_app(app_env))
    login(c, "boss", "boss-password-1")
    return c


@pytest.fixture()
def viewer(app_env):
    c = TestClient(create_app(app_env))
    login(c, "viewer1", "viewer-password-1")
    return c


@pytest.fixture()
def anon(app_env):
    return TestClient(create_app(app_env))


def conn_path(tmp_path):
    return tmp_path / "t.sqlite3"


def test_all_pages_and_api(client):
    for p in PAGES:
        r = client.get(p)
        assert r.status_code == 200, p
        assert "<html" in r.text
    for p in ["/api/health", "/api/snapshots", "/api/tables", "/api/tables/summary", "/api/tables/spare-parts",
              "/api/stores", "/api/quality", "/api/data/task?size=5"]:
        assert client.get(p).status_code == 200, p
    assert client.get("/tables/nope").status_code == 404
    assert client.get("/tables/summary?date=1999-01-01").status_code == 404
    t = client.get("/api/tables/summary").json()
    assert t["blocks"][0]["rows"][0][0] == "华北" and len(t["blocks"]) == 6


def test_filters_and_pagination(client):
    html = client.get("/tables/direct-stores", params={"q": "南宁"}).text
    assert "猛士汽车南宁江南用户中心" in html and "天津北辰" not in html
    assert client.get("/data/task", params={"q": "南宁", "size": 10}).status_code == 200
    j = client.get("/api/data/task", params={"size": 1, "page": 2}).json()
    assert j["total"] == 3 and len(j["rows"]) == 1
    j = client.get("/api/data/stock", params={"q": "22806000", "region": ""}).json()
    assert j["total"] == 2


def test_masking_by_permission(app_env, client, viewer):
    # 用户（viewer）只看汇总：明细页 / 明细接口 / 明细导出一律 403（不论 full_detail）
    for u in ("/api/data/task", "/data/task", "/export/data/task.xlsx", "/data/task?token=x"):
        assert viewer.get(u).status_code == 403, u
    assert viewer.get("/api/data/task", headers={"X-Reportdb-Token": "x"}).status_code == 403
    assert viewer.get("/").status_code == 200
    # 管理员完整
    ok = client.get("/api/data/task").json()
    assert ok["masked"] is False and ok["rows"][0]["vin"] == "LDP29H925TM016610"
    aud = TestClient(create_app(app_env))
    login(aud, "auditor", "auditor-password-1")
    assert aud.get("/api/data/task").status_code == 403


def test_interaction_feedback_markup(anon, client):
    """按钮 loading / Spin / 上传进度 / message 的样式和脚本随页面下发。"""
    login_html = anon.get("/login").text
    home = client.get("/").text
    spare = client.get("/settings/spare-parts").text
    for html in (login_html, home):
        assert "antd-feedback" in html
        assert "ant-btn-loading" in html
        assert "ant-spin-fullscreen" in html
        assert "ant-message" in html
        assert "ant-progress" in html
        assert 'id="i-loading"' in html
    assert 'type="file"' in spare and "/settings/spare-parts/upload" in spare
    assert 'class="ant-btn ant-btn-primary" type="submit"' in login_html


def test_exports(client):
    wb = load_workbook(io.BytesIO(client.get("/export/tables/all.xlsx").content))
    assert wb.sheetnames == ["总表", "区域指标", "直营店指标", "常备件备库率"]
    assert wb["直营店指标"]["A3"].value == "猛士汽车南宁江南用户中心"
    assert wb["直营店指标"]["D3"].number_format == "0.0%"
    one = load_workbook(io.BytesIO(client.get("/export/tables/spare-parts.xlsx").content))
    assert one.sheetnames == ["常备件备库率"] and one.active["F1"].value == 22806000
