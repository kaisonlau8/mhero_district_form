"""每个页面的「导出 Excel」：200、xlsx 可打开、表头、筛选生效、脱敏、权限。"""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from reportdb import exemptions as ex, store_indicators as si
from reportdb.auth import service
from reportdb.db import connect
from reportdb.webapp import create_app
from test_reportdb_exemptions import SNAP1, D1, D2  # noqa: F401
from test_reportdb_correction_exclusion_split import SHARED, load_shared
from test_reportdb_web import login  # noqa: F401

PW = {"admin": "admin-password-1", "sup1": "sup1-password-12", "view1": "view1-password-12"}


@pytest.fixture()
def env(tmp_path, conn):
    conn.execute("UPDATE stores SET supervisor='张督导' WHERE name=?", (D1,))
    conn.commit()
    load_shared(conn, SNAP1)
    service.create_user(conn, "admin", PW["admin"], role="admin")
    service.create_user(conn, "sup1", PW["sup1"], role="supervisor", supervisor_name="张督导")
    service.create_user(conn, "view1", PW["view1"], role="viewer")
    code = next(s["code"] for s in si.compute(conn, SNAP1)["stores"] if s["name"] == D1)
    conn.close()
    out = {"code": code, "path": tmp_path / "t.sqlite3"}
    for u in PW:
        c = TestClient(create_app(out["path"]))
        login(c, u, PW[u])
        out[u] = c
    return out


def wb(r):
    assert r.status_code == 200, r.text[:200]
    assert "spreadsheetml" in r.headers["content-type"]
    assert "attachment" in r.headers["content-disposition"]
    return load_workbook(BytesIO(r.content))


def heads(w, i=0):
    ws = w.worksheets[i]
    return [c.value for c in ws[1]]


def nrows(w, i=0):
    return w.worksheets[i].max_row - 1


ANY_LOGIN = ["/export/home.xlsx", "/export/quality.xlsx", "/export/tables-view/region-indicators.xlsx", "/export/tables-view/direct-stores.xlsx",
             "/export/tables-view/spare-parts.xlsx", "/export/stores.xlsx", "/export/daily-history.xlsx"]
ADMIN_ONLY = ["/export/admin/users.xlsx", "/export/admin/exclusion-rules.xlsx", "/export/admin/spare-settings.xlsx"]
REQUESTER = ["/export/exemptions/new.xlsx?indicator=second_service", "/export/exclusions/new.xlsx", "/export/exemptions/mine.xlsx"]


def test_all_exports_open_for_allowed_roles(env):
    for u in ANY_LOGIN:
        for who in ("admin", "sup1", "view1"):
            if who == "view1" and u in ("/export/quality.xlsx", "/export/stores.xlsx"):
                assert env[who].get(u).status_code == 403   # 用户只看汇总
                continue
            wb(env[who].get(u))
    for u in ADMIN_ONLY:
        wb(env["admin"].get(u))
    for u in REQUESTER:
        for who in ("admin", "sup1"):
            wb(env[who].get(u))
    wb(env["admin"].get(f"/export/stores/{env['code']}.xlsx"))
    assert env["view1"].get(f"/export/stores/{env['code']}.xlsx").status_code == 403   # 用户只看汇总


def test_permissions(env, tmp_path):
    anon = TestClient(create_app(env["path"]))
    for u in ANY_LOGIN + ADMIN_ONLY + REQUESTER + [f"/export/stores/{env['code']}.xlsx"]:
        assert anon.get(u, follow_redirects=False).status_code in (401, 303), u
    for u in ADMIN_ONLY:
        assert env["sup1"].get(u).status_code == 403, u
        assert env["view1"].get(u).status_code == 403, u
    for u in REQUESTER:
        assert env["view1"].get(u).status_code == 403, u


def test_home_export_headers_and_region_filter(env):
    w = wb(env["admin"].get("/export/home.xlsx"))
    h = heads(w)
    assert h[:2] == ["名次", "区域"] and any("招揽" in x and "完成率" in x for x in h)
    rows = [r[1].value for r in w.worksheets[0].iter_rows(min_row=2)]
    assert rows[0] == "全国" and len(rows) > 2
    reg = rows[1]
    w2 = wb(env["admin"].get("/export/home.xlsx", params={"region": reg}))
    assert [r[1].value for r in w2.worksheets[0].iter_rows(min_row=2)] == ["全国", reg]


def test_stores_export_filter_sort_and_headers(env):
    w = wb(env["admin"].get("/export/stores.xlsx"))
    h = heads(w)
    assert h[:5] == ["#", "门店", "门店编码", "区域", "直营"] and "备库率" in h
    names = [r[1].value for r in w.worksheets[0].iter_rows(min_row=2)]
    assert D1 in names and D2 in names
    w = wb(env["admin"].get("/export/stores.xlsx", params={"q": D1[:6]}))
    names = [r[1].value for r in w.worksheets[0].iter_rows(min_row=2)]
    assert D1 in names and D2 not in names
    assert nrows(wb(env["admin"].get("/export/stores.xlsx", params={"region": "不存在"}))) == 0
    wb(env["admin"].get("/export/stores.xlsx", params={"sort": "first_service"}))


def test_store_detail_export(env):
    w = wb(env["sup1"].get(f"/export/stores/{env['code']}.xlsx"))
    assert heads(w)[:3] == ["指标", "分母", "分子"] and nrows(w) == 6
    assert heads(w, 1)[:2] == ["编号", "类型"]
    assert env["sup1"].get("/export/stores/NOPE.xlsx").status_code == 404


def test_tables_view_filter(env):
    w = wb(env["admin"].get("/export/tables-view/spare-parts.xlsx"))
    assert heads(w)[:5] == ["门店", "编码", "区域", "备库率", "达标"]
    n_all = nrows(w)
    assert nrows(wb(env["admin"].get("/export/tables-view/spare-parts.xlsx", params={"q": "zzzz-none"}))) == 0 <= n_all
    assert env["admin"].get("/export/tables-view/summary.xlsx").status_code == 404


def test_apply_exports_mask_and_mine(env):
    w = wb(env["sup1"].get("/export/exemptions/new.xlsx", params={"indicator": "second_service"}))
    assert "VIN" in " ".join(str(x) for x in heads(w)) or "车架号" in " ".join(str(x) for x in heads(w))
    assert heads(w, 1)[:3] == ["编号", "类型", "VIN"]
    c = connect(env["path"])
    sup = c.execute("SELECT * FROM users WHERE username='sup1'").fetchone()
    from types import SimpleNamespace
    u = SimpleNamespace(id=sup["id"], username="sup1", role="supervisor", supervisor_name="张督导", can_request=True)
    ex.request_exemptions(c, u, "second_service", SNAP1, vins=[SHARED], category="其他", reason="补录")
    c.close()
    w = wb(env["sup1"].get("/export/exemptions/mine.xlsx"))
    assert nrows(w) == 1 and w.worksheets[0]["B2"].value == "修正结果"
    assert nrows(wb(env["sup1"].get("/export/exemptions/mine.xlsx", params={"kind": "exclude"}))) == 0
    assert nrows(wb(env["sup1"].get("/export/exemptions/mine.xlsx", params={"status": "approved"}))) == 0
    # 管理员详情导出
    w = wb(env["admin"].get("/export/admin/exemptions/1.xlsx"))
    assert heads(w)[:2] == ["编号", "类型"] and env["sup1"].get("/export/admin/exemptions/1.xlsx").status_code == 403
    assert env["admin"].get("/export/admin/exemptions/999.xlsx").status_code == 404
    wb(env["sup1"].get("/export/exclusions/new.xlsx"))


def test_admin_exports_content(env):
    w = wb(env["admin"].get("/export/admin/users.xlsx"))
    assert heads(w)[:3] == ["用户名", "姓名", "角色"]
    txt = " ".join(str(c.value) for r in w.worksheets[0].iter_rows() for c in r)
    assert "sup1" in txt and "password" not in txt.lower() and "$" not in txt   # 不含哈希
    assert heads(wb(env["admin"].get("/export/admin/exclusion-rules.xlsx")))[:2] == ["#", "数据集"]
    assert heads(wb(env["admin"].get("/export/admin/spare-settings.xlsx")), 1)[0] == "版本"
    uid = connect(env["path"]).execute("SELECT id FROM users WHERE username='sup1'").fetchone()[0]
    assert heads(wb(env["admin"].get(f"/export/admin/users/{uid}/stores.xlsx")))[:2] == ["手工", "门店"]


def test_buttons_on_pages(env):
    a = env["admin"]
    code = env["code"]
    for url, frag in [("/", "/export/home.xlsx"), ("/tables/summary", "/export/home.xlsx"), ("/tables/region-indicators", "/export/tables-view/region-indicators.xlsx"),
                      ("/tables/spare-parts", "/export/tables-view/spare-parts.xlsx"), ("/stores", "/export/stores.xlsx"), (f"/stores/{code}", f"/export/stores/{code}.xlsx"),
                      ("/exemptions/new", "/export/exemptions/new.xlsx"), ("/exclusions/new", "/export/exclusions/new.xlsx"),
                      ("/exemptions/mine", "/export/exemptions/mine.xlsx"), ("/admin/exemptions", "/admin/exemptions.xlsx"),
                      ("/admin/users", "/export/admin/users.xlsx"), ("/settings/exclusion-rules", "/export/admin/exclusion-rules.xlsx"),
                      ("/reports/daily", "/export/daily-history.xlsx"), ("/settings/spare-parts", "/export/admin/spare-settings.xlsx"),
                      ("/data/task", "/export/data/task.xlsx")]:
        r = a.get(url)
        assert r.status_code == 200 and frag in r.text and "导出 Excel" in r.text, url
    # 门店指标在「报表」菜单下
    h = a.get("/stores").text
    assert h.index("门店指标</a>") > h.index(">报表</button>") and h.index("门店指标</a>") < h.index(">明细数据</button>")
