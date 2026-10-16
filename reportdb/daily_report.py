"""系统统一导出「区域各指标情况一览MMDD.xlsx」：5 张汇总表的值（无公式）+ 可选明细 sheet。

外观沿用旧日报：以 assets/daily_report_style.xlsx（旧模板去掉明细页和全部内容后的样式骨架：列宽、字体、
填充、边框、合并单元格、条件格式）为底，把 calc 算出的值按单元格坐标写回，与旧 Excel 的版式一致。
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from . import calc
from .datasets import DATASETS
from .masking import mask_row
from .tables import Sheet

STYLE = Path(__file__).resolve().parent / "assets" / "daily_report_style.xlsx"
RANKED_VIEW = {"summary", "direct-stores"}
SUMMARY_NAMES = ["直营店指标", "区域指标排序", "直营店指标排序", "总表", "常备件备库率"]


def report_filename(snap: str) -> str:
    """区域各指标情况一览MMDD.xlsx（与旧日报同名规则）。"""
    d = datetime.strptime(snap, "%Y-%m-%d")
    return f"区域各指标情况一览{d.strftime('%m%d')}.xlsx"


def build_daily_report(conn, snap: str, *, include_detail: bool = False, full_detail: bool = False) -> bytes:
    sheets = calc.compute(conn, snap)
    wb = load_workbook(STYLE)
    by_name = {sh.name: sh for sh in sheets.values()}
    for ws in wb.worksheets:
        sh = by_name.get(ws.title)
        if sh is None:
            continue
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c, MergedCell) and c.value is not None:
                    c.value = None
        if sh.slug in RANKED_VIEW:   # 总表/直营店指标：每个指标块按名次 1..N 排（只改行顺序，值不变；与「排序」页同序）
            sh = Sheet(sh.slug, sh.name, [b.sorted_copy() for b in sh.blocks], sh.extra_cells, sh.extra_fmts, sh.matrix)
        fm = sh.formats()
        for (r, c), v in sh.grid().items():
            cell = ws.cell(r, c)
            if isinstance(cell, MergedCell):
                continue
            cell.value = v
            if fm.get((r, c)):
                cell.number_format = fm[(r, c)]
    if include_detail:
        for ds in DATASETS:
            ws = wb.create_sheet(ds.template_sheet)
            cols = [c for c, _ in ds.columns]
            ws.append([lab for _, lab in ds.columns])
            where = " AND in_scope=1" if ds.has_region else ""
            for r in conn.execute(f"SELECT {','.join(cols)} FROM {ds.table} WHERE snapshot_date=?{where} ORDER BY row_no", (snap,)):
                d = mask_row(ds, dict(r), full_detail)
                ws.append([d[c] for c in cols])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_daily_report(conn, snap: str, output_dir: Path, **kw) -> Path:
    from datetime import datetime as _dt

    output_dir.mkdir(parents=True, exist_ok=True)
    out = output_dir / report_filename(snap)
    if out.exists():  # 与旧流水线一致：同日重复生成不覆盖
        out = output_dir / f"{out.stem}_{_dt.now().strftime('%H%M%S')}.xlsx"
    out.write_bytes(build_daily_report(conn, snap, **kw))
    return out
