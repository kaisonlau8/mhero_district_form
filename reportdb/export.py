"""xlsx 导出（只写值，格式用 0% / 0.0%，不含公式）。"""
from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .tables import Sheet


def sheets_to_xlsx(sheets: Iterable[Sheet]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(sh.name)
        for r, c, v, fmt in sh.cells():
            cell = ws.cell(r, c, v)
            if fmt:
                cell.number_format = fmt
        for b in sh.blocks:
            ws.cell(b.row, b.col).font = Font(bold=True)
        if sh.matrix is not None:
            for c in range(1, 6):
                ws.cell(1, c).font = Font(bold=True)
            ws.freeze_panes = "F5"
            ws.column_dimensions["C"].width = 34
        else:
            for c in range(1, 16):
                ws.column_dimensions[get_column_letter(c)].width = 22 if c in (1, 6, 11) else 11
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rows_to_xlsx(title: str, headers: list[str], rows: Iterable[list[Any]]) -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def multi_rows_to_xlsx(sheets: list[tuple[str, list[str], Iterable[list[Any]]]]) -> bytes:
    """多个 sheet：(sheet 名, 表头, 行)。表头加粗、冻结首行、按内容估算列宽；只写值。"""
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for title, headers, rows in sheets:
        base = "".join("_" if ch in '[]:*?/\\' else ch for ch in (title or "Sheet"))[:31] or "Sheet"
        name, n = base, 2
        while name in used:
            name = f"{base[:28]}_{n}"
            n += 1
        used.add(name)
        ws = wb.create_sheet(name)
        ws.append(list(headers))
        widths = [min(40, max(8, len(str(h)) * 2)) for h in headers]
        for r in rows:
            r = ["" if v is None else v for v in r]
            ws.append(r)
            for i, v in enumerate(r[:len(widths)]):
                widths[i] = min(40, max(widths[i], len(str(v)) + 2))
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
