"""对账工具：用「Excel 公式引擎」（模板 + app/excel_values.py 纯 Python 求值）当参照，逐格对比 reportdb 的计算结果。

参照构造方式（与生产 Excel 路径一致）：
  1. 取现行模板 app/assets/report_template.xlsx；
  2. 把源数据按 processor.py 同样的方式写入 备件库存明细/招揽实施率/首保/二保/新保/续保；
  3. app.excel_values 把公式算成值；
  4. 排序页的 SORT() 数组公式 excel_values 不支持，这里按公式文本（SORT(range,5,1)）独立实现一次。
region_map 不为 None 时，先把不在门店表内的明细行整行剔除，再把 5 张明细 sheet 的 B 列（大区）改成按门店映射后的区域，
用于核对「按门店表归属区域」模式（即模板公式 COUNTIF(B:B,"华北") 想要的输入）。
"""
from __future__ import annotations

import re
import sys
import tempfile
from pathlib import Path
from typing import Any, Callable

from . import config
from .calc import SLUGS
from .datasets import BY_KEY, DATASETS

if str(config.PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(config.PROJECT_ROOT))

SHEET_NAMES = list(SLUGS.values())
LAST_FORMULA_XLSX: Path | None = None  # 最近一次构造的「含公式」参照工作簿，可交给 LibreOffice 做第三方复算
_SORT_RE = re.compile(r"SORT\(\s*(?:(?P<sheet>[^!(),]+)!)?(?P<rng>[A-Z]+\d+:[A-Z]+\d+)\s*,\s*(?P<col>\d+)\s*,\s*(?P<asc>-?1)\s*\)")


def _coords(rng: str):
    from openpyxl.utils import range_boundaries

    c1, r1, c2, r2 = range_boundaries(rng)
    return r1, c1, r2, c2


def reference_values(data: dict[str, list[list[Any]]], region_map: Callable[[Any, Any], Any] | None = None,
                     workdir: Path | None = None) -> dict[str, dict[tuple[int, int], Any]]:
    """data: {dataset_key: 对齐/归一化后的行}（loaders 输出或日报 xlsx 里的源 sheet）。返回 {sheet名: {(r,c): 值}}。"""
    from openpyxl import load_workbook
    from openpyxl.worksheet.formula import ArrayFormula

    from app.excel_values import flatten_formulas_to_values

    wb = load_workbook(config.TEMPLATE_PATH, data_only=False)
    for ds in DATASETS:
        ws = wb[ds.template_sheet]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        names = [c for c, _ in ds.columns]
        ci = names.index(ds.store_code_col) if ds.store_code_col else None
        ni = names.index(ds.store_name_col) if ds.store_name_col else None
        for row in data[ds.key]:
            row = list(row) + [None] * (len(names) - len(row))
            row = row[: len(names)]
            row = [None if v == "" else v for v in row]
            if region_map is not None and ds.has_region:
                reg, scope = region_map(row[ci], row[ni])
                if not scope:  # 门店不在门店表（多维表格）内：整行剔除，不计入任何指标（含直营店、全国）
                    continue
                row[1] = reg
            ws.append(row)
    tmp = Path(tempfile.mkdtemp(prefix="rdb_ref_", dir=workdir))
    src = tmp / "ref_formulas.xlsx"
    wb.save(src)
    global LAST_FORMULA_XLSX
    LAST_FORMULA_XLSX = src
    arrays: dict[str, list[tuple[str, str, str]]] = {}
    for name in ("区域指标排序", "直营店指标排序"):
        arrays[name] = []
        for row in wb[name].iter_rows():
            for c in row:
                if isinstance(c.value, ArrayFormula):
                    arrays[name].append((c.coordinate, c.value.ref, c.value.text))
    wb.close()
    out_path = flatten_formulas_to_values(src, tmp / "ref_values.xlsx")
    vwb = load_workbook(out_path, data_only=True)
    grids: dict[str, dict[tuple[int, int], Any]] = {}
    for name in SHEET_NAMES:
        ws = vwb[name]
        grids[name] = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    vwb.close()
    # SORT() 数组公式：excel_values 不求值，这里按公式文本独立实现
    for name, items in arrays.items():
        g = grids[name]
        for _, ref, text in items:
            m = _SORT_RE.search(text)
            assert m, f"无法解析 SORT 公式: {text}"
            src_name = m.group("sheet")
            r1, c1, r2, c2 = _coords(m.group("rng"))
            sg = grids[src_name]
            rows = [[sg.get((r, c)) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]
            k = int(m.group("col")) - 1
            rows.sort(key=lambda row: (row[k] is None, row[k] if row[k] is not None else 0), reverse=(m.group("asc") == "-1"))
            tr1, tc1, _, _ = _coords(ref)
            for i, row in enumerate(rows):
                for j, v in enumerate(row):
                    if v is None:
                        g.pop((tr1 + i, tc1 + j), None)
                    else:
                        g[(tr1 + i, tc1 + j)] = v
    # 排序页里模板残留的旧缓存值（数组区域之外的普通引用单元格已被 excel_values 重算，无需处理）
    return grids


def _num(v: Any):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def compare(mine: dict[str, Any], ref: dict[str, dict[tuple[int, int], Any]], tol: float = 1e-9) -> dict[str, Any]:
    """mine: {slug: Sheet}。返回对比统计与差异列表。"""
    from .tables import col_letter

    total = exact = approx = 0
    max_delta = 0.0
    approx_cells: list[dict[str, Any]] = []
    diffs: list[dict[str, Any]] = []
    per_sheet: dict[str, int] = {}
    for slug, sheet in mine.items():
        g = sheet.grid()
        rg = ref[sheet.name]
        keys = set(g) | set(rg)
        n = 0
        for k in sorted(keys):
            a, b = g.get(k), rg.get(k)
            if a is None and b in (None, ""):
                continue
            n += 1
            na, nb = _num(a), _num(b)
            if na is not None and nb is not None:
                if na == nb:
                    exact += 1
                elif abs(na - nb) <= tol:
                    approx += 1
                    max_delta = max(max_delta, abs(na - nb))
                    approx_cells.append({'sheet': sheet.name, 'cell': f'{col_letter(k[1])}{k[0]}', 'mine': a, 'ref': b})
                else:
                    diffs.append({"sheet": sheet.name, "cell": f"{col_letter(k[1])}{k[0]}", "mine": a, "ref": b})
            elif a == b or (isinstance(a, str) and isinstance(b, str) and a.strip() == b.strip()):
                exact += 1
            else:
                diffs.append({"sheet": sheet.name, "cell": f"{col_letter(k[1])}{k[0]}", "mine": a, "ref": b})
        per_sheet[sheet.name] = n
        total += n
    return {"cells": total, "exact": exact, "approx": approx, "max_float_delta": max_delta, "approx_cells": approx_cells, "diff_count": len(diffs), "diffs": diffs, "per_sheet": per_sheet}


def main(argv: list[str] | None = None) -> int:
    import argparse
    import json

    from .calc import compute
    from .db import connect, migrate
    from .ingest import ingest_dir, ingest_report
    from .loaders import load_directory_sources, load_report_workbook
    from .stores import region_lookup_factory

    ap = argparse.ArgumentParser(description="reportdb 与 Excel 公式引擎逐格对账（临时库，不影响正式库）")
    ap.add_argument("--dir")
    ap.add_argument("--report")
    ap.add_argument("--mode", choices=["store", "source"], default="store")
    args = ap.parse_args(argv)
    tmpdb = Path(tempfile.mkdtemp(prefix="rdb_parity_")) / "p.sqlite3"
    conn = connect(tmpdb)
    migrate(conn)
    if args.dir:
        data, _ = load_directory_sources(Path(args.dir))
        data.pop("_origin_renewal")
        ingest_dir(conn, Path(args.dir), "2000-01-01", recompute=False)
    else:
        data = load_report_workbook(Path(args.report))
        ingest_report(conn, Path(args.report), "2000-01-01", recompute=False)
    mine = compute(conn, "2000-01-01", args.mode)
    ref = reference_values(data, region_lookup_factory(conn) if args.mode == "store" else None)
    res = compare(mine, ref)
    print(json.dumps({k: v for k, v in res.items() if k != "diffs"}, ensure_ascii=False))
    for d in res["diffs"][:50]:
        print(d)
    return 0 if res["diff_count"] == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
