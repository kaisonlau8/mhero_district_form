"""常备件名单：版本化、上传解析、预览（匹配库存）、生效、回滚、模板文件。"""
from __future__ import annotations

import csv
import io
import json
import re
import sqlite3
from datetime import datetime
from typing import Any

from .loaders import part_key

HEAD_CODE, HEAD_NAME, HEAD_COUNTED, HEAD_STD, HEAD_GROUP = "备件号", "备件名称", "是否计入分母", "标准备库量", "分组"
_YES = {"是", "y", "yes", "true", "1", "计入", "√", "✓"}
_NO = {"否", "n", "no", "false", "0", "不计入", "×", "x"}


class SpareListError(ValueError):
    pass


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ----------------------------------------------------------------------------- 读取
def list_for_date(conn: sqlite3.Connection, snap: str) -> dict[str, Any]:
    """快照日期当天生效的名单：effective_from <= 日期 中最新的一版（同日取最后创建的）；都晚于该日期则取最早的一版。"""
    row = conn.execute("SELECT * FROM spare_lists WHERE effective_from<=? ORDER BY effective_from DESC, id DESC LIMIT 1",
                       (snap,)).fetchone() or conn.execute("SELECT * FROM spare_lists ORDER BY effective_from, id LIMIT 1").fetchone()
    if not row:
        raise SpareListError("没有常备件名单")
    return get_list(conn, row["id"])


def get_list(conn: sqlite3.Connection, list_id: int) -> dict[str, Any]:
    head = conn.execute("SELECT * FROM spare_lists WHERE id=?", (list_id,)).fetchone()
    if not head:
        raise SpareListError("名单版本不存在")
    items = [dict(r) for r in conn.execute("SELECT * FROM spare_list_items WHERE list_id=? ORDER BY sort_order", (list_id,))]
    return {**dict(head), "items": items}


def history(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    out = []
    for r in conn.execute("SELECT * FROM spare_lists ORDER BY effective_from DESC, id DESC"):
        n, c = conn.execute("SELECT COUNT(*), COALESCE(SUM(counted),0) FROM spare_list_items WHERE list_id=?", (r["id"],)).fetchone()
        out.append({**dict(r), "n_items": n, "n_counted": c})
    return out


# ----------------------------------------------------------------------------- 解析上传
def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise SpareListError("无法识别 CSV 编码，请另存为 UTF-8")


def _cell_text(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v == int(v):
        v = int(v)
    s = str(v).strip()
    return s or None


def parse_upload(filename: str, raw: bytes) -> list[dict[str, Any]]:
    """xlsx / csv → [{code,name,counted,std_qty,group_label}]。表头需含「备件号」「备件名称」。"""
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        rows = list(csv.reader(io.StringIO(_decode(raw))))
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise SpareListError(f"无法读取 xlsx：{exc}") from exc
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        raise SpareListError("只支持 .xlsx 或 .csv 文件")
    rows = [r for r in rows if any(_cell_text(c) for c in r)]
    if not rows:
        raise SpareListError("文件是空的")
    head = [(_cell_text(c) or "").replace(" ", "") for c in rows[0]]
    for need in (HEAD_CODE, HEAD_NAME):
        if need not in head:
            raise SpareListError(f"表头缺少「{need}」列（需要：备件号、备件名称，可选：是否计入分母、标准备库量、分组）")
    ix = {h: head.index(h) for h in (HEAD_CODE, HEAD_NAME, HEAD_COUNTED, HEAD_STD, HEAD_GROUP) if h in head}
    out = []
    for n, r in enumerate(rows[1:], 2):
        def get(h):
            i = ix.get(h)
            return _cell_text(r[i]) if i is not None and i < len(r) else None
        code = get(HEAD_CODE)
        if not code:
            continue
        flag = (get(HEAD_COUNTED) or "是").lower()
        if flag in _YES:
            counted = 1
        elif flag in _NO:
            counted = 0
        else:
            raise SpareListError(f"第 {n} 行「是否计入分母」只能填 是/否（或 1/0），当前为「{flag}」")
        std = get(HEAD_STD)
        try:
            std_v = float(std) if std is not None else None
        except ValueError:
            raise SpareListError(f"第 {n} 行「标准备库量」不是数字：{std}")
        out.append({"code": code, "name": get(HEAD_NAME), "counted": counted, "std_qty": std_v, "group_label": get(HEAD_GROUP)})
    if not out:
        raise SpareListError("没有读到任何备件号")
    return out


# ----------------------------------------------------------------------------- 预览
def preview(conn: sqlite3.Connection, parts: list[dict[str, Any]], snap: str | None) -> dict[str, Any]:
    """与最近一个快照的库存比对：每个备件是否匹配、库存行数、有货店数；整体覆盖率与试算结果。"""
    from . import calc  # 避免循环

    warnings: list[str] = []
    seen: dict[str, int] = {}
    for i, p in enumerate(parts):
        k = part_key(p["code"])
        p["key"] = k
        if k in seen:
            warnings.append(f"备件号重复：{p['code']}（与第 {seen[k] + 1} 条规范化后相同）")
        else:
            seen[k] = i
    if not any(p["counted"] for p in parts):
        raise SpareListError("至少需要一个「计入分母」的备件")
    info: dict[str, Any] = {"snapshot_date": snap, "parts": parts, "warnings": warnings}
    if not snap:
        info.update(rows=parts, matched=None, unmatched=None, coverage=None, trial=None)
        return info
    rows_by_key = {r[0]: (r[1], r[2]) for r in conn.execute(
        "SELECT part_key, COUNT(*), COUNT(DISTINCT store_name) FROM stock_items WHERE snapshot_date=? GROUP BY part_key", (snap,))}
    spare = calc.spare_matrix(conn, snap, parts=parts)
    store_n = len(spare["rows"])
    have = {}
    for j, p in enumerate(parts):
        have[j] = sum(1 for s in spare["rows"] if s["vals"][j] > 0)
    matched = unmatched = 0
    rows = []
    for j, p in enumerate(parts):
        st = rows_by_key.get(p["key"])
        ok = st is not None
        matched += ok
        unmatched += not ok
        rows.append({**p, "matched": ok, "stock_rows": st[0] if st else 0, "stores_with_stock": have[j], "store_total": store_n})
    thr = float(conn.execute("SELECT value FROM meta WHERE key='spare_threshold'").fetchone()[0] or 0.8)
    cur = calc.spare_matrix(conn, snap)
    def summarize(sp):
        ok = sum(1 for s in sp["rows"] if s["rate"] is not None and s["rate"] >= thr)
        avg = sum(s["rate"] or 0 for s in sp["rows"]) / len(sp["rows"]) if sp["rows"] else None
        return {"stores": len(sp["rows"]), "ok_stores": ok, "avg_rate": avg, "denominator": sp["counted"]}
    info.update(rows=rows, matched=matched, unmatched=unmatched, coverage=matched / len(parts),
                trial=summarize(spare), current=summarize(cur))
    return info


def save_draft(conn: sqlite3.Connection, user: str, filename: str, parts: list[dict[str, Any]]) -> int:
    with conn:
        cur = conn.execute("INSERT INTO spare_drafts(created_at,created_by,filename,payload) VALUES (?,?,?,?)",
                           (_now(), user, filename, json.dumps(parts, ensure_ascii=False)))
    return cur.lastrowid


def load_draft(conn: sqlite3.Connection, draft_id: int) -> dict[str, Any]:
    r = conn.execute("SELECT * FROM spare_drafts WHERE id=?", (draft_id,)).fetchone()
    if not r:
        raise SpareListError("预览已过期，请重新上传")
    return {**dict(r), "parts": json.loads(r["payload"])}


# ----------------------------------------------------------------------------- 生效 / 回滚
def create_version(conn: sqlite3.Connection, parts: list[dict[str, Any]], effective_from: str, user: str, *,
                   note: str = "", source: str = "upload", filename: str = "") -> int:
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", effective_from or ""):
        raise SpareListError("生效日期需为 YYYY-MM-DD")
    if not parts or not any(p["counted"] for p in parts):
        raise SpareListError("至少需要一个「计入分母」的备件")
    with conn:
        cur = conn.execute("INSERT INTO spare_lists(effective_from,created_at,created_by,note,source,filename) VALUES (?,?,?,?,?,?)",
                           (effective_from, _now(), user, note, source, filename))
        lid = cur.lastrowid
        conn.executemany(
            "INSERT INTO spare_list_items(list_id,sort_order,code,name,std_qty,group_label,counted) VALUES (?,?,?,?,?,?,?)",
            [(lid, i, p["code"], p.get("name"), p.get("std_qty"), p.get("group_label"), 1 if p.get("counted", 1) else 0)
             for i, p in enumerate(parts, 1)])
    return lid


def rollback(conn: sqlite3.Connection, list_id: int, effective_from: str, user: str) -> int:
    old = get_list(conn, list_id)
    return create_version(conn, old["items"], effective_from, user, note=f"回滚到 v{list_id}",
                          source=f"rollback:v{list_id}")


def recompute_from(conn: sqlite3.Connection, effective_from: str) -> list[str]:
    """名单变更后：重算 effective_from 及之后的所有快照（汇总表按当天生效的名单）。"""
    from . import calc

    done = []
    for r in conn.execute("SELECT snapshot_date FROM snapshots WHERE snapshot_date>=? ORDER BY snapshot_date", (effective_from,)).fetchall():
        calc.compute_and_store(conn, r[0])
        done.append(r[0])
    return done


# ----------------------------------------------------------------------------- 模板文件
def template_xlsx(items: list[dict[str, Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "常备件名单"
    ws.append([HEAD_CODE, HEAD_NAME, HEAD_COUNTED, HEAD_STD, HEAD_GROUP])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8F0FF")
    for it in items:
        q = it.get("std_qty")
        ws.append([it["code"], it.get("name"), "是" if it.get("counted", 1) else "否",
                   (int(q) if q is not None and float(q) == int(q) else q), it.get("group_label")])
    for col, w in zip("ABCDE", (16, 34, 14, 12, 10)):
        ws.column_dimensions[col].width = w
    for r in range(2, ws.max_row + 1):  # 备件号一律文本，避免 Excel 把 00123 / 55003001WC 改写
        ws.cell(r, 1).number_format = "@"
    h = wb.create_sheet("填写说明")
    for line in ("1. 必填列：备件号、备件名称。",
                 "2. 「是否计入分母」可不填，默认「是」；填「否」表示该备件仍会统计库存，但不计入备库率的分母/分子（如 817 保护套件）。",
                 "3. 「标准备库量」「分组」只用于展示，不参与计算。",
                 "4. 匹配规则：备件号去掉空格/横线、不分大小写；数字和文本形式的数字视为相同。",
                 "5. 上传后先预览匹配结果，确认后才生效。"):
        h.append([line])
    h.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
